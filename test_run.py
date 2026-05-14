import os
import sys
from docx import Document

# Add src to path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), 'src')))

from navidoc.core import NaviDoc

def create_test_files():
    os.makedirs("storage/test_data", exist_ok=True)
    
    # 1. Create Markdown
    md_file = "storage/test_data/test.md"
    if not os.path.exists(md_file):
        print(f"Creating test file at {md_file}...")
        with open(md_file, 'w', encoding='utf-8') as f:
            f.write("# Header 1\nThis is the content of section 1.\n\n## Header 2\nThis is the secret code in MD: 12345.\n")

    # 2. Create DOCX
    docx_file = "storage/test_data/test.docx"
    if not os.path.exists(docx_file):
        print(f"Creating test file at {docx_file}...")
        doc = Document()
        doc.add_heading('DOCX Header 1', level=1)
        doc.add_paragraph('This is section 1 content in DOCX.')
        doc.add_heading('DOCX Header 2', level=2)
        doc.add_paragraph('The secret code in DOCX is 54321.')
        doc.save(docx_file)
        
    return md_file, docx_file

def main():
    print("Initializing NaviDoc...")
    engine = NaviDoc(model="phi3") 
    
    md_file, docx_file = create_test_files()
    
    # Test Markdown
    print(f"\nIngesting {md_file}...")
    print(engine.ingest(md_file))
    
    query = "What is the secret code in Header 2?"
    print(f"Querying: '{query}'")
    try:
        print("Response:", engine.query(query))
    except Exception as e:
        print(f"Error: {e}")

    # Test DOCX
    print(f"\nIngesting {docx_file}...")
    print(engine.ingest(docx_file))
    
    query = "What is the secret code in DOCX Header 2?"
    print(f"Querying: '{query}'")
    try:
        print("Response:", engine.query(query))
    except Exception as e:
        print(f"Error: {e}")

if __name__ == "__main__":
    main()
