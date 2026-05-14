from pptx import Presentation
from typing import Dict, Any, List

class PptxParser:
    def __init__(self):
        pass

    def parse(self, file_path: str) -> Dict[str, Any]:
        """Parse a PPTX file."""
        prs = Presentation(file_path)
        slides = []
        
        for i, slide in enumerate(prs.slides):
            text = []
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text.append(shape.text)
            slides.append({
                "page_number": i + 1,
                "content": "\n".join(text)
            })
            
        return {
            "title": file_path,
            "type": "pptx",
            "pages": slides
        }
